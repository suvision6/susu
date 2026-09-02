#!/usr/bin/env python3
"""Deterministic human-facing XLSX renderer for director-shot-data/3.1.3.

The spreadsheet is the human frontend.  JSON, Markdown and validation retain
the richer Agent-facing backend; this writer projects the same facts through
``shot_flow`` instead of copying backend ``execution_text`` into column five.
"""

from __future__ import annotations

import argparse
import json
import math
import os
import re
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional

try:  # supports script-path and ``python -m scripts.export_xlsx``
    from .storyboard_delivery import load_json, normalize_text, safe_slug, validate_delivery
    from ._xlsx_projection import (
        project_xlsx,
        projected_dialogue_segment_count,
        projection_sha256,
    )
except ImportError:  # pragma: no cover - script-path execution
    from storyboard_delivery import load_json, normalize_text, safe_slug, validate_delivery
    from _xlsx_projection import (  # type: ignore[no-redef]
        project_xlsx,
        projected_dialogue_segment_count,
        projection_sha256,
    )

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception as exc:  # pragma: no cover - minimal runtimes
    Workbook = None  # type: ignore[assignment]
    load_workbook = None  # type: ignore[assignment]
    OPENPYXL_ERROR: Optional[Exception] = exc
else:
    OPENPYXL_ERROR = None

HEADERS = ("镜号", "场景", "原剧本段落", "镜头时长", "运镜＋主画面描述", "备注")
XLSX_METADATA_TIME = datetime(2000, 1, 1, 0, 0, 0)
XLSX_ZIP_TIME = (2000, 1, 1, 0, 0, 0)


def bounded_row_height(row: list[Any]) -> float:
    """Estimate a readable wrapped height without creating unbounded rows."""
    widths = (8, 18, 38, 10, 70, 24)
    line_counts: list[int] = []
    for value, width in zip(row, widths):
        text = normalize_text(value)
        logical_lines = text.split("\n") if text else [""]
        count = sum(max(1, math.ceil(len(line) / max(width, 1))) for line in logical_lines)
        line_counts.append(count)
    return float(min(180, max(28, 16 * max(line_counts) + 8)))


def _storyboard_rows(data: dict[str, Any]) -> list[list[Any]]:
    """Compatibility helper returning the canonical XLSX projection rows."""
    return project_xlsx(data)["storyboard_rows"]


def build_openpyxl_workbook(data: dict[str, Any]):
    if OPENPYXL_ERROR is not None or Workbook is None:
        raise RuntimeError(f"openpyxl 不可用：{OPENPYXL_ERROR}")

    wb = Workbook()
    # openpyxl otherwise stamps each build with the wall clock, which makes
    # the artifact hash change even when the canonical shot model is unchanged.
    wb.properties.creator = "su-fenjingskill"
    wb.properties.lastModifiedBy = "su-fenjingskill"
    wb.properties.created = XLSX_METADATA_TIME
    wb.properties.modified = XLSX_METADATA_TIME
    storyboard = wb.active
    storyboard.title = "导演分镜"
    design_sheet = wb.create_sheet("导演设计")

    projection = project_xlsx(data)
    title = projection["title"]
    rows = projection["storyboard_rows"]

    title_fill = PatternFill("solid", fgColor="111827")
    subtitle_fill = PatternFill("solid", fgColor="E5E7EB")
    header_fill = PatternFill("solid", fgColor="374151")
    cjk_font = "Hiragino Sans GB"
    title_font = Font(name=cjk_font, bold=True, color="FFFFFF", size=16)
    header_font = Font(name=cjk_font, bold=True, color="FFFFFF")
    body_font = Font(name=cjk_font, size=10, color="111827")
    body_alignment = Alignment(vertical="top", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    storyboard.merge_cells("A1:F1")
    storyboard["A1"] = f"{title}｜导演分镜"
    storyboard["A1"].fill = title_fill
    storyboard["A1"].font = title_font
    storyboard["A1"].alignment = Alignment(horizontal="left", vertical="center")

    storyboard.merge_cells("A2:F2")
    storyboard["A2"] = projection["summary"]
    storyboard["A2"].fill = subtitle_fill
    storyboard["A2"].font = body_font
    storyboard["A2"].alignment = Alignment(horizontal="left", vertical="center")

    for column, value in enumerate(HEADERS, start=1):
        cell = storyboard.cell(row=4, column=column, value=value)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment

    for row in rows or [["", "", "", "", "", ""]]:
        storyboard.append(row)

    for row_number in range(5, storyboard.max_row + 1):
        for column in range(1, 7):
            cell = storyboard.cell(row=row_number, column=column)
            cell.alignment = body_alignment
            cell.font = body_font
        storyboard.cell(row=row_number, column=1).alignment = center_alignment
        storyboard.cell(row=row_number, column=4).alignment = center_alignment
        storyboard.cell(row=row_number, column=4).number_format = '0.0"秒"'
        source_row = rows[row_number - 5] if row_number - 5 < len(rows) else [""] * 6
        storyboard.row_dimensions[row_number].height = bounded_row_height(source_row)

    for index, width in enumerate((9, 18, 34, 11, 64, 22), start=1):
        storyboard.column_dimensions[get_column_letter(index)].width = width
    storyboard.row_dimensions[1].height = 30
    storyboard.row_dimensions[2].height = 22
    storyboard.row_dimensions[4].height = 26
    storyboard.freeze_panes = "A5"
    storyboard.auto_filter.ref = f"A4:F{storyboard.max_row}"
    storyboard.sheet_view.showGridLines = False
    storyboard.sheet_properties.pageSetUpPr.fitToPage = True
    storyboard.page_setup.orientation = "landscape"
    storyboard.page_setup.paperSize = storyboard.PAPERSIZE_A4
    storyboard.page_setup.fitToWidth = 1
    storyboard.page_setup.fitToHeight = 0
    storyboard.print_title_rows = "1:4"
    storyboard.print_area = f"A1:F{storyboard.max_row}"
    storyboard.page_margins.left = 0.2
    storyboard.page_margins.right = 0.2
    storyboard.page_margins.top = 0.3
    storyboard.page_margins.bottom = 0.3

    design_sheet.merge_cells("A1:B1")
    design_sheet["A1"] = f"{title}｜导演设计摘要"
    design_sheet["A1"].fill = title_fill
    design_sheet["A1"].font = title_font
    design_sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    design_sheet["A3"] = "维度"
    design_sheet["B3"] = "导演设计"
    for cell in design_sheet[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment

    for row in projection["director_design_rows"]:
        design_sheet.append(row)

    assumption_header_row = design_sheet.max_row + 2
    design_sheet.merge_cells(
        start_row=assumption_header_row,
        start_column=1,
        end_row=assumption_header_row,
        end_column=2,
    )
    header_cell = design_sheet.cell(row=assumption_header_row, column=1, value="假设与待确认项")
    header_cell.fill = header_fill
    header_cell.font = header_font
    header_cell.alignment = center_alignment
    for row in projection["assumption_rows"]:
        design_sheet.append(row)

    for row in design_sheet.iter_rows(min_row=4, max_row=design_sheet.max_row, min_col=1, max_col=2):
        for cell in row:
            if cell.row == assumption_header_row:
                continue
            cell.alignment = body_alignment
            cell.font = body_font
    design_sheet.column_dimensions["A"].width = 18
    design_sheet.column_dimensions["B"].width = 88
    design_sheet.row_dimensions[1].height = 30
    for row_number in range(4, design_sheet.max_row + 1):
        if row_number == assumption_header_row:
            design_sheet.row_dimensions[row_number].height = 26
            continue
        value = normalize_text(design_sheet.cell(row=row_number, column=2).value)
        lines = max(1, sum(max(1, math.ceil(len(line) / 76)) for line in value.split("\n")))
        design_sheet.row_dimensions[row_number].height = float(min(160, max(28, 16 * lines + 8)))
    design_sheet.freeze_panes = "A4"
    design_sheet.sheet_view.showGridLines = False
    design_sheet.sheet_properties.pageSetUpPr.fitToPage = True
    design_sheet.page_setup.orientation = "portrait"
    design_sheet.page_setup.paperSize = design_sheet.PAPERSIZE_A4
    design_sheet.page_setup.fitToWidth = 1
    design_sheet.page_setup.fitToHeight = 0
    design_sheet.print_title_rows = "1:3"
    design_sheet.print_area = f"A1:B{design_sheet.max_row}"
    design_sheet.page_margins.left = 0.3
    design_sheet.page_margins.right = 0.3

    if wb.sheetnames != ["导演分镜", "导演设计"]:
        raise RuntimeError(f"正式 XLSX 工作表漂移：{wb.sheetnames}")
    return wb


# Compatibility alias. There is only one backend.
def build_workbook(data: dict[str, Any]):
    return build_openpyxl_workbook(data)


def _normalize_xlsx_archive(path: Path) -> None:
    """Normalize ZIP member timestamps so identical workbooks hash identically."""
    with zipfile.ZipFile(path, "r") as source:
        members = []
        for item in source.infolist():
            payload = source.read(item.filename)
            if item.filename == "docProps/core.xml":
                payload = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<2>",
                    payload,
                    count=1,
                )
            members.append((item, payload))

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        for original, payload in members:
            normalized = zipfile.ZipInfo(original.filename, date_time=XLSX_ZIP_TIME)
            normalized.compress_type = zipfile.ZIP_DEFLATED
            normalized.create_system = 0
            normalized.external_attr = original.external_attr
            normalized.internal_attr = original.internal_attr
            normalized.comment = original.comment
            target.writestr(normalized, payload)


def export_xlsx(data: dict[str, Any], output: Path) -> Path:
    """Write one XLSX atomically and refuse any overwrite."""
    if output.exists():
        raise FileExistsError(f"拒绝覆盖已有 XLSX：{output}")
    if OPENPYXL_ERROR is not None or Workbook is None:
        raise RuntimeError(f"openpyxl 不可用：{OPENPYXL_ERROR}")
    output.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output.stem}-",
        suffix=".tmp.xlsx",
        dir=str(output.parent),
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        wb = build_openpyxl_workbook(data)
        wb.save(temporary)
        _normalize_xlsx_archive(temporary)
        verify_xlsx(data, temporary)
        os.replace(temporary, output)
    finally:
        if temporary.exists():
            temporary.unlink()
    return output


def _cell_value_for_compare(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, str):
        return normalize_text(value)
    return value


def verify_xlsx(data: dict[str, Any], path: Path) -> dict[str, Any]:
    """Verify both human-facing sheets against their canonical projection."""
    if OPENPYXL_ERROR is not None or load_workbook is None:
        raise RuntimeError(f"openpyxl 不可用：{OPENPYXL_ERROR}")
    projection = project_xlsx(data)
    wb = load_workbook(path, data_only=False, read_only=False)
    try:
        if wb.sheetnames != ["导演分镜", "导演设计"]:
            raise ValueError(f"XLSX 工作表不一致：{wb.sheetnames}")
        sheet = wb["导演分镜"]
        expected_title = f"{projection['title']}｜导演分镜"
        if _cell_value_for_compare(sheet["A1"].value) != expected_title:
            raise ValueError(f"XLSX 标题不一致：{sheet['A1'].value!r}")
        if _cell_value_for_compare(sheet["A2"].value) != projection["summary"]:
            raise ValueError(f"XLSX 摘要不一致：{sheet['A2'].value!r}")
        actual_headers = tuple(sheet.cell(row=4, column=i).value for i in range(1, 7))
        if actual_headers != HEADERS:
            raise ValueError(f"XLSX 六列表头不一致：{actual_headers}")
        expected_rows = projection["storyboard_rows"]
        actual_count = max(0, sheet.max_row - 4)
        if actual_count != len(expected_rows):
            raise ValueError(f"XLSX 镜头行数不一致：expected={len(expected_rows)}, actual={actual_count}")
        for row_index, expected in enumerate(expected_rows, start=5):
            actual = [sheet.cell(row=row_index, column=column).value for column in range(1, 7)]
            for column_index in (0, 1, 2, 3, 5):
                if _cell_value_for_compare(actual[column_index]) != _cell_value_for_compare(expected[column_index]):
                    raise ValueError(
                        f"XLSX 第 {row_index} 行第 {column_index + 1} 列与后端事实不一致："
                        f"expected={expected[column_index]!r}, actual={actual[column_index]!r}"
                    )
            if _cell_value_for_compare(actual[4]) != _cell_value_for_compare(expected[4]):
                raise ValueError(
                    f"XLSX 第 {row_index} 行第五列与 canonical XLSX projection 不一致："
                    f"expected={expected[4]!r}, actual={actual[4]!r}"
                )
        if sheet.freeze_panes != "A5":
            raise ValueError(f"XLSX 冻结窗格漂移：{sheet.freeze_panes!r}")
        if sheet.auto_filter.ref != f"A4:F{sheet.max_row}":
            raise ValueError(f"XLSX 筛选区域漂移：{sheet.auto_filter.ref!r}")

        design_sheet = wb["导演设计"]
        expected_design_title = f"{projection['title']}｜导演设计摘要"
        if _cell_value_for_compare(design_sheet["A1"].value) != expected_design_title:
            raise ValueError(f"导演设计标题不一致：{design_sheet['A1'].value!r}")
        if (design_sheet["A3"].value, design_sheet["B3"].value) != ("维度", "导演设计"):
            raise ValueError("导演设计表头不一致。")
        design_rows = projection["director_design_rows"]
        for row_index, expected in enumerate(design_rows, start=4):
            actual = [design_sheet.cell(row=row_index, column=column).value for column in (1, 2)]
            if [_cell_value_for_compare(value) for value in actual] != [
                _cell_value_for_compare(value) for value in expected
            ]:
                raise ValueError(
                    f"导演设计第 {row_index} 行不一致：expected={expected!r}, actual={actual!r}"
                )
        assumption_header_row = 4 + len(design_rows) + 1
        if design_sheet.cell(row=assumption_header_row, column=1).value != "假设与待确认项":
            raise ValueError("导演设计缺少假设与待确认项标题。")
        assumption_rows = projection["assumption_rows"]
        for row_index, expected in enumerate(assumption_rows, start=assumption_header_row + 1):
            actual = [design_sheet.cell(row=row_index, column=column).value for column in (1, 2)]
            if [_cell_value_for_compare(value) for value in actual] != [
                _cell_value_for_compare(value) for value in expected
            ]:
                raise ValueError(
                    f"待确认项第 {row_index} 行不一致：expected={expected!r}, actual={actual!r}"
                )
        expected_design_max_row = assumption_header_row + len(assumption_rows)
        if design_sheet.max_row != expected_design_max_row:
            raise ValueError(
                f"导演设计行数不一致：expected={expected_design_max_row}, actual={design_sheet.max_row}"
            )
        if design_sheet.freeze_panes != "A4":
            raise ValueError(f"导演设计冻结窗格漂移：{design_sheet.freeze_panes!r}")

        return {
            "status": "PASS",
            "sheet_names": wb.sheetnames,
            "sheet_parity": {"导演分镜": "PASS", "导演设计": "PASS"},
            "shot_row_count": len(expected_rows),
            "projection_sha256": projection_sha256(data),
            "dialogue_segment_count": projected_dialogue_segment_count(data),
            "director_design_field_count": len(design_rows),
        }
    finally:
        wb.close()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="将通过 combined gate 的 director-shot-data/3.1.3 导出为 XLSX。")
    parser.add_argument("--input", type=Path, required=True, help="输入 JSON 文件。")
    parser.add_argument("--workspace", type=Path, required=True, help="director-workspace/3.1.3。")
    parser.add_argument("--output", type=Path, help="输出 XLSX；省略时按 delivery_slug 命名。")
    parser.add_argument("--fail-on-warn", action="store_true", help="将 READY_WITH_ASSUMPTIONS 视为失败。")
    return parser


def main(argv: Optional[Iterable[str]] = None) -> int:
    args = build_parser().parse_args(list(argv) if argv is not None else None)
    if OPENPYXL_ERROR is not None:
        sys.stderr.write(f"FAIL: openpyxl 不可用：{OPENPYXL_ERROR}\n")
        return 2
    try:
        data = load_json(args.input)
        workspace = load_json(args.workspace)
        report, materialized = validate_delivery(data, workspace)
        if report["status"] == "FAIL" or materialized is None:
            sys.stderr.write(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
            return 1
        if report["status"] == "READY_WITH_ASSUMPTIONS" and args.fail_on_warn:
            sys.stderr.write(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
            return 1
        output = args.output or Path(f"{safe_slug(materialized)}-storyboard.xlsx")
        export_xlsx(materialized, output)
    except (OSError, ValueError, RuntimeError) as exc:
        sys.stderr.write(f"FAIL: XLSX 导出失败：{exc}\n")
        return 2
    sys.stdout.write(f"xlsx: {output}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
