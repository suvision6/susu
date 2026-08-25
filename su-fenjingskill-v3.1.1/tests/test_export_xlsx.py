#!/usr/bin/env python3

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from export_xlsx import OPENPYXL_ERROR, build_openpyxl_workbook, export_xlsx


@unittest.skipIf(OPENPYXL_ERROR is not None, "openpyxl unavailable")
class ExportXlsxTests(unittest.TestCase):
    def setUp(self) -> None:
        self.data = json.loads(
            (ROOT / "examples" / "kitchen-farewell-shot-data.json").read_text(
                encoding="utf-8"
            )
        )

    def test_workbook_has_exactly_two_frozen_sheets(self) -> None:
        workbook = build_openpyxl_workbook(self.data)
        self.assertEqual(["导演分镜", "导演设计"], workbook.sheetnames)
        self.assertEqual("A5", workbook["导演分镜"].freeze_panes)
        self.assertEqual("A4", workbook["导演设计"].freeze_panes)

    def test_primary_headers_are_frozen_contract(self) -> None:
        workbook = build_openpyxl_workbook(self.data)
        sheet = workbook["导演分镜"]
        headers = [sheet.cell(row=4, column=index).value for index in range(1, 7)]
        self.assertEqual(
            ["镜号", "场景", "原剧本段落", "镜头时长", "运镜＋主画面描述", "备注"],
            headers,
        )

    def test_export_roundtrip_keeps_only_two_sheets(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "storyboard.xlsx"
            export_xlsx(self.data, output)
            workbook = load_workbook(output, read_only=False, data_only=True)
            self.assertEqual(["导演分镜", "导演设计"], workbook.sheetnames)
            self.assertEqual(6, workbook["导演分镜"].max_column)

    def test_production_risk_is_on_design_sheet_not_remarks_column(self) -> None:
        risk = "桌面反光与钥匙焦点需要现场联调。"
        self.data["shots"][0]["production_risks"] = [risk]
        self.data["shots"][0]["notes"] = ""
        workbook = build_openpyxl_workbook(self.data)
        storyboard = workbook["导演分镜"]
        design = workbook["导演设计"]
        self.assertEqual("", storyboard["F5"].value or "")
        design_values = [
            cell.value
            for row in design.iter_rows()
            for cell in row
            if cell.value is not None
        ]
        self.assertIn("制作风险（不进入备注列）", design_values)
        self.assertIn(risk, design_values)
        self.assertEqual("Hiragino Sans GB", storyboard["E5"].font.name)


if __name__ == "__main__":
    unittest.main()
