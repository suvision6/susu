#!/usr/bin/env python3

from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from export_xlsx import OPENPYXL_ERROR, build_openpyxl_workbook, export_xlsx, verify_xlsx  # noqa: E402
from _xlsx_projection import (  # noqa: E402
    XlsxProjectionError,
    build_dialogue_index,
    render_xlsx_execution_text,
)


@unittest.skipIf(OPENPYXL_ERROR is not None, "openpyxl unavailable")
class ExportXlsxTests(unittest.TestCase):
    def setUp(self) -> None:
        self.data = json.loads((ROOT / "examples" / "kitchen-farewell-shot-data.json").read_text(encoding="utf-8"))
        self.data["shots"][0]["shot_flow"] = [
            {"owner": "camera_setup"},
            {"owner": "blocking"},
            {"owner": "effect", "index": 0},
            {"owner": "effect", "index": 1},
            {"owner": "performance"},
            {"owner": "ambience"},
            {"owner": "state_update", "index": 0},
            {"owner": "edit_exit"},
        ]
        self.data["shots"][1]["shot_flow"] = [
            {"owner": "camera_setup"},
            {"owner": "blocking"},
            {"owner": "dialogue_segment", "index": 0},
            {"owner": "performance"},
            {"owner": "effect", "index": 0},
            {"owner": "effect", "index": 1},
            {"owner": "ambience"},
            {"owner": "state_update", "index": 0},
            {"owner": "edit_exit"},
        ]
        self.data["shots"][2]["shot_flow"] = [
            {"owner": "camera_setup"},
            {"owner": "blocking"},
            {"owner": "performance"},
            {"owner": "dialogue_segment", "index": 0},
            {"owner": "ambience"},
            {"owner": "state_update", "index": 0},
            {"owner": "edit_exit"},
        ]

    def test_workbook_has_exactly_two_frozen_sheets(self) -> None:
        workbook = build_openpyxl_workbook(self.data)
        self.assertEqual(["导演分镜", "导演设计"], workbook.sheetnames)
        self.assertEqual("A5", workbook["导演分镜"].freeze_panes)
        self.assertEqual("A4", workbook["导演设计"].freeze_panes)

    def test_primary_headers_are_frozen_contract(self) -> None:
        workbook = build_openpyxl_workbook(self.data)
        sheet = workbook["导演分镜"]
        headers = [sheet.cell(row=4, column=index).value for index in range(1, 7)]
        self.assertEqual(["镜号", "场景", "原剧本段落", "镜头时长", "运镜＋主画面描述", "备注"], headers)
        self.assertEqual("A4:F7", sheet.auto_filter.ref)
        self.assertIn("$A$1:$F$7", str(sheet.print_area))
        self.assertEqual("画幅 16:9｜节奏 自定义节奏｜对白 自定义 4.5字/秒（区间 4–5）｜共 3 镜｜总时长 14.5 秒", sheet["A2"].value)
        self.assertNotIn("contract", sheet["A2"].value)
        self.assertNotIn("project", sheet["A2"].value)

    def test_export_roundtrip_and_parity(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / "storyboard.xlsx"
            export_xlsx(self.data, output)
            workbook = load_workbook(output, read_only=False, data_only=True)
            self.assertEqual(["导演分镜", "导演设计"], workbook.sheetnames)
            self.assertEqual(6, workbook["导演分镜"].max_column)
            verification = verify_xlsx(self.data, output)
            self.assertEqual("PASS", verification["status"])
            self.assertEqual({"导演分镜": "PASS", "导演设计": "PASS"}, verification["sheet_parity"])
            self.assertEqual(2, verification["dialogue_segment_count"])
            self.assertEqual(10, verification["director_design_field_count"])
            self.assertRegex(verification["projection_sha256"], r"^[0-9a-f]{64}$")


    def test_identical_inputs_produce_byte_identical_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            first = Path(directory) / "first.xlsx"
            second = Path(directory) / "second.xlsx"
            export_xlsx(self.data, first)
            export_xlsx(self.data, second)
            self.assertEqual(first.read_bytes(), second.read_bytes())

    def test_export_refuses_to_overwrite_existing_xlsx(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "historical-storyboard.xlsx"
            sentinel = b"historical-xlsx-bytes"
            output.write_bytes(sentinel)
            with self.assertRaises(FileExistsError):
                export_xlsx(self.data, output)
            self.assertEqual(sentinel, output.read_bytes())

    def test_notes_are_rendered_only_in_the_sixth_column(self) -> None:
        data = copy.deepcopy(self.data)
        note = "车窗反射与同期收音需由现场部门自行处理。"
        data["shots"][0]["notes"] = note
        workbook = build_openpyxl_workbook(data)
        storyboard = workbook["导演分镜"]
        design = workbook["导演设计"]
        self.assertEqual(note, storyboard["F5"].value)
        design_values = "\n".join(
            str(cell.value)
            for row in design.iter_rows()
            for cell in row
            if cell.value is not None
        )
        self.assertNotIn(note, design_values)
        self.assertEqual("Hiragino Sans GB", storyboard["E5"].font.name)

    def test_fifth_column_is_compact_timeline_projection_not_execution_text(self) -> None:
        data = copy.deepcopy(self.data)
        data["shots"][1]["execution_text"] = "【摄影】D001 Gate Alignment 音乐：无"
        workbook = build_openpyxl_workbook(data)
        text = workbook["导演分镜"]["E6"].value
        self.assertEqual(1, text.count("\n"))
        self.assertTrue(text.startswith("【微俯视，近景，固定】\n【画面内容】"))
        self.assertIn("林晓彤在画外说：“我明天走。”", text)
        for forbidden in (
            "【摄影】",
            "【调度与表演】",
            "【声音】",
            "【剪辑】",
            "【连续性】",
            "【时长】",
            "【镜头动机】",
            "D001",
            "Gate",
            "Alignment",
            "音乐：无",
            "来源动作按序",
        ):
            self.assertNotIn(forbidden, text)

    def test_flow_spans_are_exact_backend_substrings_and_keep_time_order(self) -> None:
        shot = copy.deepcopy(self.data["shots"][2])
        shot["shot_flow"] = [
            {"owner": "camera_setup"},
            {"owner": "blocking", "span": "不转身"},
            {"owner": "dialogue_segment", "index": 0},
            {"owner": "edit_exit", "span": "仍看向画外左侧"},
        ]
        text = render_xlsx_execution_text(shot, build_dialogue_index(self.data))
        self.assertLess(text.index("不转身"), text.index("陈默说"))
        self.assertLess(text.index("陈默说"), text.index("仍看向画外左侧"))
        broken = copy.deepcopy(shot)
        broken["shot_flow"][1]["span"] = "转身面对她"
        with self.assertRaisesRegex(XlsxProjectionError, "逐字片段"):
            render_xlsx_execution_text(broken, build_dialogue_index(self.data))

    def test_dialogue_flow_must_be_complete_unique_and_ordered(self) -> None:
        shot = copy.deepcopy(self.data["shots"][1])
        shot["sound"]["dialogue_segments"].append(
            {"dialogue_id": "D001", "text": "我明天走。", "delivery": "os"}
        )
        shot["shot_flow"] = [
            {"owner": "camera_setup"},
            {"owner": "dialogue_segment", "index": 1},
            {"owner": "dialogue_segment", "index": 0},
        ]
        with self.assertRaisesRegex(XlsxProjectionError, "不重不漏并保持顺序"):
            render_xlsx_execution_text(shot, build_dialogue_index(self.data))

    def test_nonfixed_movement_is_required_once_and_renders_without_reason(self) -> None:
        shot = copy.deepcopy(self.data["shots"][0])
        shot["camera"]["movement"] = {
            "type": "push",
            "trigger": "钥匙触桌",
            "speed": "缓慢",
            "path": "从双人关系推进到两人之间的钥匙",
            "end_condition": "钥匙在桌面停稳",
            "reason": "后端保留的运动理由不进入前端",
        }
        with self.assertRaisesRegex(XlsxProjectionError, "恰好引用一次 movement"):
            render_xlsx_execution_text(shot, build_dialogue_index(self.data))
        shot["shot_flow"].insert(4, {"owner": "movement"})
        text = render_xlsx_execution_text(shot, build_dialogue_index(self.data))
        self.assertTrue(text.startswith("【平视，中远景，推进】"))
        for phrase in ("钥匙触桌", "缓慢", "从双人关系推进到两人之间的钥匙", "钥匙在桌面停稳"):
            self.assertIn(phrase, text)
        self.assertNotIn("后端保留的运动理由", text)

    def test_fixed_shot_must_not_reference_movement(self) -> None:
        shot = copy.deepcopy(self.data["shots"][0])
        shot["shot_flow"].insert(1, {"owner": "movement"})
        with self.assertRaisesRegex(XlsxProjectionError, "固定镜头不得"):
            render_xlsx_execution_text(shot, build_dialogue_index(self.data))

    def test_camera_setup_owner_shape_and_flow_indices_are_strict(self) -> None:
        shot = copy.deepcopy(self.data["shots"][0])
        shot["shot_flow"][0]["span"] = "厨房入口"
        with self.assertRaisesRegex(XlsxProjectionError, "不允许的字段"):
            render_xlsx_execution_text(shot, build_dialogue_index(self.data))
        shot = copy.deepcopy(self.data["shots"][0])
        shot["shot_flow"][2]["index"] = 99
        with self.assertRaisesRegex(XlsxProjectionError, "越界"):
            render_xlsx_execution_text(shot, build_dialogue_index(self.data))

    def test_design_sheet_has_ten_summary_dimensions_and_natural_assumptions(self) -> None:
        data = copy.deepcopy(self.data)
        data["assumptions"] = [
            {
                "assumption_id": "A001",
                "scope": "SC001",
                "statement": "门位于餐桌左后方。",
                "reason": "来源未说明门的位置。",
                "impact": "只影响人物离开时的银幕方向。",
                "status": "open",
            }
        ]
        sheet = build_openpyxl_workbook(data)["导演设计"]
        labels = [sheet.cell(row=row, column=1).value for row in range(4, 14)]
        self.assertEqual(10, len(labels))
        self.assertEqual("场景任务", labels[0])
        self.assertEqual("节奏策略", labels[-1])
        self.assertEqual("假设与待确认项", sheet["A15"].value)
        self.assertEqual("待确认项 1", sheet["A16"].value)
        self.assertIn("门位于餐桌左后方。", sheet["B16"].value)
        self.assertIn("影响：只影响人物离开时的银幕方向。", sheet["B16"].value)
        self.assertNotIn("A001", sheet["A16"].value)
        self.assertNotIn("open", sheet["A16"].value)

    def test_verify_rejects_tampered_fifth_column_and_design_sheet(self) -> None:
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as directory:
            original = Path(directory) / "original.xlsx"
            export_xlsx(self.data, original)
            workbook = load_workbook(original)
            workbook["导演分镜"]["E5"] = "被人为改写"
            fifth_tampered = Path(directory) / "fifth-tampered.xlsx"
            workbook.save(fifth_tampered)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "canonical XLSX projection"):
                verify_xlsx(self.data, fifth_tampered)

            workbook = load_workbook(original)
            workbook["导演设计"]["B4"] = "被人为改写"
            design_tampered = Path(directory) / "design-tampered.xlsx"
            workbook.save(design_tampered)
            workbook.close()
            with self.assertRaisesRegex(ValueError, "导演设计第"):
                verify_xlsx(self.data, design_tampered)


if __name__ == "__main__":
    unittest.main()
