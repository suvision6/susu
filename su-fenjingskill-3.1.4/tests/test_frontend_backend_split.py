#!/usr/bin/env python3

from __future__ import annotations

import copy
import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

from _execution_text import canonical_execution_text  # noqa: E402
from _xlsx_projection import (  # noqa: E402
    build_dialogue_index,
    project_xlsx,
    projection_sha256,
    render_xlsx_execution_text,
)
from export_xlsx import (  # noqa: E402
    OPENPYXL_ERROR,
    build_openpyxl_workbook,
    export_xlsx,
    verify_xlsx,
)
from migrate_contract_3_1_2_to_3_1_3 import build_drafts  # noqa: E402
from storyboard_delivery import (  # noqa: E402
    build_outputs,
    markdown_cell,
    validate_structure,
)
from storyboard_review import validate_shot_flow, validate_workspace  # noqa: E402


FORBIDDEN_FRONTEND_TEXT = (
    "【摄影】",
    "【调度与表演】",
    "【声音】",
    "【剪辑】",
    "【连续性】",
    "【时长】",
    "【镜头动机】",
    "音乐：无",
    "状态变化：无",
    "未定义",
    "来源动作按序",
    "Gate",
    "Alignment",
    "以人物共享关系为主体",
    "自然透视，保留人物真实距离",
    "焦点优先保持人物关系可读",
    "保持当前场景客观声场",
    "关系轴同侧",
    "保持到当前关系动作与迟到反应完整落地",
)
MACHINE_ID_RE = re.compile(r"(?<![A-Za-z0-9])(?:D|F|P|A|G|RF)[0-9]{3,}(?![A-Za-z0-9])")
FRONTEND_RE = re.compile(r"^【[^【】\n]+，[^【】\n]+，[^【】\n]+】\n【画面内容】[^\n]+$")


def error_codes(report: dict) -> set[str]:
    return {item["code"] for item in report.get("errors", [])}


class FrontendBackendContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.example = json.loads(
            (ROOT / "examples" / "kitchen-farewell-shot-data.json").read_text(encoding="utf-8")
        )
        cls.workspace = json.loads(
            (ROOT / "examples" / "kitchen-farewell-director-workspace.json").read_text(encoding="utf-8")
        )
        cls.concept = json.loads(
            (ROOT / "examples" / "unknown-room-awakening-shot-data.json").read_text(encoding="utf-8")
        )

    def test_contract_identity_is_3_1_4(self) -> None:
        self.assertEqual("3.1.4", self.example["contract_version"])
        self.assertEqual("3.1.4", self.example["source_skill_version"])
        self.assertEqual("director-workspace/3.1.4", self.workspace["workspace_contract"])

        formal_schema = json.loads(
            (ROOT / "schemas" / "director-shot-data.schema.json").read_text(encoding="utf-8")
        )
        workspace_schema = json.loads(
            (ROOT / "schemas" / "director-workspace.schema.json").read_text(encoding="utf-8")
        )
        self.assertEqual("director-shot-data/3.1.4", formal_schema["title"])
        self.assertEqual("3.1.4", formal_schema["properties"]["contract_version"]["const"])
        self.assertEqual("3.1.4", formal_schema["properties"]["source_skill_version"]["const"])
        self.assertEqual("director-workspace/3.1.4", workspace_schema["title"])

    def test_nonstandard_shot_size_term_is_blocking(self) -> None:
        draft = copy.deepcopy(self.example)
        draft["shots"][0]["camera"]["shot_size"] = "紧中景"
        draft["shots"][0]["execution_text"] = canonical_execution_text(draft["shots"][0])
        self.assertIn("SHOT_SIZE_TERM_NONSTANDARD", error_codes(validate_structure(draft)))

    def test_canonical_shot_size_transition_remains_valid(self) -> None:
        draft = copy.deepcopy(self.example)
        draft["shots"][0]["camera"]["shot_size"] = "中景→近景"
        draft["shots"][0]["execution_text"] = canonical_execution_text(draft["shots"][0])
        self.assertNotIn("SHOT_SIZE_TERM_NONSTANDARD", error_codes(validate_structure(draft)))

    def test_camera_composition_cannot_repeat_flow_action(self) -> None:
        draft = copy.deepcopy(self.example)
        shot = draft["shots"][0]
        shot["camera"]["composition"] = shot["staging"]["blocking"]
        shot["execution_text"] = canonical_execution_text(shot)
        self.assertIn("XLSX_PROJECTION_REDUNDANT_CONTENT", error_codes(validate_structure(draft)))

    def test_camera_position_cannot_repeat_header_terms(self) -> None:
        draft = copy.deepcopy(self.example)
        shot = draft["shots"][0]
        shot["camera"]["position"] = "平视中远景，位于厨房入口内侧"
        shot["execution_text"] = canonical_execution_text(shot)
        self.assertIn("XLSX_CAMERA_SETUP_REDUNDANT", error_codes(validate_structure(draft)))

    def test_generic_focus_should_stay_out_of_xlsx_flow(self) -> None:
        draft = copy.deepcopy(self.example)
        shot = draft["shots"][0]
        shot["camera"]["focus"] = "焦点停在两人，到当前反应落定。"
        shot["execution_text"] = canonical_execution_text(shot)
        self.assertIn("XLSX_PROJECTION_NONESSENTIAL_FOCUS", error_codes(validate_structure(draft)))

    def test_repeated_optional_ambience_across_scene_is_blocking(self) -> None:
        draft = copy.deepcopy(self.example)
        repeated = draft["shots"][0]["sound"]["ambience"]
        draft["shots"][1]["sound"]["ambience"] = repeated
        draft["shots"][1]["execution_text"] = canonical_execution_text(draft["shots"][1])
        self.assertIn("XLSX_PROJECTION_REPEATED_OPTIONAL_FLOW", error_codes(validate_structure(draft)))

    def test_camera_output_evals_assert_real_json_result_structure(self) -> None:
        rows = [
            json.loads(line)
            for line in (ROOT / "evals" / "output" / "cases.jsonl").read_text(encoding="utf-8").splitlines()
            if line.strip()
        ]
        structured = [row for row in rows if row.get("metadata", {}).get("case_type") == "structured_result"]
        self.assertGreaterEqual(len(structured), 4)

        def resolve(value, path: str):
            current = value
            for name, index_text in re.findall(r"([^.\[\]]+)(?:\[(\d+)\])?", path):
                current = current[name]
                if index_text:
                    current = current[int(index_text)]
            return current

        for row in structured:
            self.assertFalse(row["metadata"]["static_keywords_sufficient"])
            self.assertEqual("json", row["execution"]["result_format"])
            result = json.loads(row["with_skill_output"])
            checks = row.get("structural_assertions", [])
            self.assertTrue(checks, row["id"])
            for check in checks:
                actual = resolve(result, check["path"])
                if "equals" in check:
                    self.assertEqual(check["equals"], actual, f"{row['id']}:{check['path']}")
                if "contains" in check:
                    self.assertIn(check["contains"], actual, f"{row['id']}:{check['path']}")
                if check.get("not_empty"):
                    self.assertTrue(actual, f"{row['id']}:{check['path']}")

    def test_valid_examples_have_complete_shot_flow_and_n_minus_one_boundaries(self) -> None:
        self.assertEqual([], validate_shot_flow(self.example["shots"]))
        for strategy in self.workspace["scene_strategies"]:
            topology = strategy["topology"]
            self.assertGreater(len(topology), 0)
            self.assertEqual(
                max(0, len(topology) - 1),
                sum("boundary_to_next" in unit for unit in topology),
            )
            for unit in topology[:-1]:
                boundary = unit["boundary_to_next"]
                self.assertTrue(boundary["relation"].strip())
                self.assertTrue(boundary["trigger"].strip())
                self.assertTrue(boundary["editorial_gain"].strip())
            self.assertNotIn("boundary_to_next", topology[-1])

    def test_shot_flow_rejects_empty_owner_index_span_and_camera_count_errors(self) -> None:
        empty = copy.deepcopy(self.example["shots"])
        empty[0]["shot_flow"] = []
        self.assertIn("SHOT_FLOW_REQUIRED", {item["code"] for item in validate_shot_flow(empty)})

        bad_owner = copy.deepcopy(self.example["shots"])
        bad_owner[0]["shot_flow"][0]["owner"] = "risk_register"
        self.assertIn("SHOT_FLOW_OWNER_INVALID", {item["code"] for item in validate_shot_flow(bad_owner)})

        bad_index = copy.deepcopy(self.example["shots"])
        bad_index[0]["shot_flow"].append({"owner": "effect", "index": 999})
        self.assertIn("SHOT_FLOW_INDEX_OUT_OF_RANGE", {item["code"] for item in validate_shot_flow(bad_index)})

        bad_span = copy.deepcopy(self.example["shots"])
        bad_span[0]["shot_flow"].append({"owner": "blocking", "span": "后端字段里不存在的自由正文"})
        self.assertIn("SHOT_FLOW_SPAN_MISMATCH", {item["code"] for item in validate_shot_flow(bad_span)})

        duplicate_camera = copy.deepcopy(self.example["shots"])
        duplicate_camera[0]["shot_flow"].append({"owner": "camera_setup"})
        self.assertIn(
            "SHOT_FLOW_CAMERA_SETUP_COUNT",
            {item["code"] for item in validate_shot_flow(duplicate_camera)},
        )

    def test_shot_flow_rejects_dialogue_loss_duplication_and_reordering(self) -> None:
        shots = copy.deepcopy(self.example["shots"])
        dialogue_shot = next(shot for shot in shots if shot["sound"]["dialogue_segments"])
        dialogue_shot["shot_flow"] = [
            item for item in dialogue_shot["shot_flow"] if item["owner"] != "dialogue_segment"
        ]
        codes = {item["code"] for item in validate_shot_flow([dialogue_shot])}
        self.assertIn("SHOT_FLOW_DIALOGUE_MISSING", codes)
        self.assertIn("SHOT_FLOW_DIALOGUE_ORDER", codes)

        duplicate = copy.deepcopy(self.example["shots"])
        dialogue_shot = next(shot for shot in duplicate if shot["sound"]["dialogue_segments"])
        insertion = next(
            index for index, item in enumerate(dialogue_shot["shot_flow"]) if item["owner"] == "dialogue_segment"
        )
        dialogue_shot["shot_flow"].insert(insertion + 1, {"owner": "dialogue_segment", "index": 0})
        codes = {item["code"] for item in validate_shot_flow([dialogue_shot])}
        self.assertIn("SHOT_FLOW_DIALOGUE_DUPLICATE", codes)
        self.assertIn("SHOT_FLOW_DIALOGUE_ORDER", codes)

        reordered = copy.deepcopy(self.example["shots"])
        first = next(shot for shot in reordered if shot["sound"]["dialogue_segments"])
        second = next(
            shot for shot in reordered
            if shot is not first and shot["sound"]["dialogue_segments"]
        )
        first["sound"]["dialogue_segments"].append(copy.deepcopy(second["sound"]["dialogue_segments"][0]))
        dialogue_positions = [
            index for index, item in enumerate(first["shot_flow"]) if item["owner"] == "dialogue_segment"
        ]
        insert_at = dialogue_positions[-1]
        first["shot_flow"][insert_at] = {"owner": "dialogue_segment", "index": 1}
        first["shot_flow"].insert(insert_at + 1, {"owner": "dialogue_segment", "index": 0})
        self.assertIn(
            "SHOT_FLOW_DIALOGUE_ORDER",
            {item["code"] for item in validate_shot_flow([first])},
        )

    def test_shot_flow_enforces_real_movement_and_omits_fixed_movement(self) -> None:
        moving = copy.deepcopy(self.example["shots"])
        moving[0]["camera"]["movement"] = {
            "type": "push",
            "trigger": "钥匙触桌",
            "speed": "缓慢地",
            "path": "沿两人之间的负空间向餐桌推进",
            "end_condition": "钥匙在桌面停稳",
            "reason": "让关系断裂在物件落地后变得可见。",
        }
        self.assertIn(
            "SHOT_FLOW_MOVEMENT_REQUIRED",
            {item["code"] for item in validate_shot_flow(moving)},
        )

        fixed = copy.deepcopy(self.example["shots"])
        fixed[0]["shot_flow"].append({"owner": "movement"})
        self.assertIn(
            "SHOT_FLOW_FIXED_MOVEMENT_FORBIDDEN",
            {item["code"] for item in validate_shot_flow(fixed)},
        )

    def test_workspace_rejects_missing_vague_and_terminal_boundaries(self) -> None:
        missing = copy.deepcopy(self.workspace)
        missing["scene_strategies"][0]["topology"][0].pop("boundary_to_next")
        report = validate_workspace(missing, self.example, check_review_states=False)
        self.assertIn("BOUNDARY_TO_NEXT_MISSING", error_codes(report))

        vague = copy.deepcopy(self.workspace)
        vague_boundary = vague["scene_strategies"][0]["topology"][0]["boundary_to_next"]
        vague_boundary["trigger"] = "动作变化"
        vague_boundary["editorial_gain"] = "增强节奏"
        report = validate_workspace(vague, self.example, check_review_states=False)
        self.assertIn("BOUNDARY_TRIGGER_VAGUE", error_codes(report))
        self.assertIn("BOUNDARY_EDITORIAL_GAIN_VAGUE", error_codes(report))

        terminal = copy.deepcopy(self.workspace)
        terminal["scene_strategies"][0]["topology"][-1]["boundary_to_next"] = {
            "relation": "cut",
            "trigger": "场景已经结束",
            "editorial_gain": "切到不存在的下一镜。",
        }
        report = validate_workspace(terminal, self.example, check_review_states=False)
        self.assertIn("TERMINAL_BOUNDARY_FORBIDDEN", error_codes(report))

    def test_warning_does_not_masquerade_as_open_assumption(self) -> None:
        warning_only = copy.deepcopy(self.example)
        warning_only["source"]["delivery_slug"] = "中文临时名"
        report = validate_structure(warning_only)
        self.assertEqual([], report["errors"])
        self.assertGreater(report["warning_count"], 0)
        self.assertEqual("READY", report["status"])

        concept_report = validate_structure(self.concept)
        self.assertEqual([], concept_report["errors"])
        self.assertGreater(concept_report["summary"]["open_assumption_count"], 0)
        self.assertEqual("READY_WITH_ASSUMPTIONS", concept_report["status"])

    def test_contract_has_no_generic_duration_text_or_shot_count_ceiling(self) -> None:
        schema = json.loads(
            (ROOT / "schemas" / "director-shot-data.schema.json").read_text(encoding="utf-8")
        )
        self.assertNotIn("maxItems", schema["properties"]["shots"])
        self.assertNotIn("maximum", schema["$defs"]["shot"]["properties"]["duration_seconds"])
        self.assertNotIn("maxLength", schema["$defs"]["nonEmptyString"])

        long_take = copy.deepcopy(self.example)
        long_take["shots"][0]["duration_seconds"] = 3600
        long_take["shots"][0]["duration_basis"] = "按完整实时过程保留一小时，不因通用镜长阈值拆开。"
        long_take["shots"][0]["timing_plan"]["blocks"][0]["hold_seconds"] = 3600
        long_take["shots"][0]["timing_plan"]["blocks"][0]["computed_seconds"] = 3600
        long_take["shots"][0]["timing_plan"]["computed_duration_seconds"] = 3600
        long_take["shots"][0]["staging"]["blocking"] += "随后保持等待。" * 1200
        long_take["shots"][0]["execution_text"] = canonical_execution_text(long_take["shots"][0])
        report = validate_structure(long_take)
        self.assertEqual([], report["errors"])
        self.assertEqual("READY", report["status"])

    def test_312_migration_creates_unapproved_rebuild_drafts_without_guessing(self) -> None:
        legacy_data = copy.deepcopy(self.example)
        legacy_data["contract_version"] = "3.1.2"
        legacy_data["source_skill_version"] = "3.1.2"
        for shot in legacy_data["shots"]:
            shot["edit"]["reason"] = "旧版重复切点证明。"
            shot["motivation"]["cut_or_hold_reason"] = "旧版重复留切证明。"
            shot.pop("shot_flow", None)

        legacy_workspace = copy.deepcopy(self.workspace)
        legacy_workspace["workspace_contract"] = "director-workspace/3.1.2"
        for strategy in legacy_workspace["scene_strategies"]:
            for index, unit in enumerate(strategy["topology"]):
                unit.pop("boundary_to_next", None)
                unit["inter_shot_relation"] = "scene_end" if index == len(strategy["topology"]) - 1 else "cut"
                unit["boundary_reason"] = {
                    "source_change": "旧来源变化",
                    "mechanism_need": "旧机制说明",
                    "method_basis": "旧方法依据",
                    "alternative_rejected": "旧替代方案",
                }

        shot_draft, workspace_draft, report = build_drafts(legacy_data, legacy_workspace)
        self.assertEqual("3.1.3", shot_draft["contract_version"])
        self.assertTrue(all(shot["shot_flow"] == [] for shot in shot_draft["shots"]))
        self.assertTrue(all("reason" not in shot["edit"] for shot in shot_draft["shots"]))
        self.assertTrue(
            all("cut_or_hold_reason" not in shot["motivation"] for shot in shot_draft["shots"])
        )
        self.assertEqual("director-workspace/3.1.3", workspace_draft["workspace_contract"])
        self.assertEqual("invalidated", workspace_draft["review_lock"]["gate_1_status"])
        self.assertEqual("invalidated", workspace_draft["review_lock"]["gate_2_status"])
        self.assertEqual("invalidated", workspace_draft["review_lock"]["alignment_status"])
        self.assertEqual("DRAFT_REVIEW_REQUIRED", report["status"])
        self.assertFalse(report["formal_ready"])

    @staticmethod
    def _formal_codes(data: dict) -> set[str]:
        for shot in data["shots"]:
            shot["execution_text"] = canonical_execution_text(shot)
        return {item["code"] for item in validate_structure(data)["errors"]}

    def _workspace_codes(self, workspace: dict, data: dict) -> set[str]:
        return error_codes(validate_workspace(workspace, data, check_review_states=False))

    def test_schema_exposes_every_framing_mode_without_a_category_quota(self) -> None:
        schema = json.loads(
            (ROOT / "schemas" / "director-shot-data.schema.json").read_text(encoding="utf-8")
        )
        modes = set(schema["$defs"]["camera"]["properties"]["framing_mode"]["enum"])
        self.assertTrue(
            {"single", "two_shot", "group", "over_shoulder", "insert", "subjective", "space"}
            <= modes
        )
        self.assertNotIn("minContains", schema["properties"]["shots"])

    def test_visibility_partition_and_dialogue_delivery_are_enforced(self) -> None:
        overlap = copy.deepcopy(self.example)
        shot = overlap["shots"][0]
        subject = shot["staging"]["visible_subjects"][0]
        shot["staging"]["offscreen_subjects"].append(subject)
        self.assertIn("SUBJECT_VISIBILITY_OVERLAP", self._formal_codes(overlap))

        undeclared = copy.deepcopy(self.example)
        shot = undeclared["shots"][0]
        shot["staging"]["visible_subjects"] = []
        shot["staging"]["offscreen_subjects"] = []
        self.assertIn("SUBJECT_VISIBILITY_UNDECLARED", self._formal_codes(undeclared))

        source_speakers = {
            item["dialogue_id"]: item["speaker"] for item in self.example["source"]["dialogue_lines"]
        }
        for delivery, target_bucket, other_bucket in (
            ("os", "visible_subjects", "offscreen_subjects"),
            ("onscreen", "offscreen_subjects", "visible_subjects"),
        ):
            changed = copy.deepcopy(self.example)
            dialogue_shot = next(
                shot
                for shot in changed["shots"]
                if any(item["delivery"] == delivery for item in shot["sound"]["dialogue_segments"])
            )
            segment = next(
                item for item in dialogue_shot["sound"]["dialogue_segments"] if item["delivery"] == delivery
            )
            speaker = source_speakers[segment["dialogue_id"]]
            for bucket in ("visible_subjects", "offscreen_subjects"):
                dialogue_shot["staging"][bucket] = [
                    item for item in dialogue_shot["staging"][bucket] if item != speaker
                ]
            dialogue_shot["staging"][target_bucket].append(speaker)
            self.assertIn("DIALOGUE_VISIBILITY_MISMATCH", self._formal_codes(changed))

    def test_framing_cardinality_over_shoulder_and_insert_semantics(self) -> None:
        for mode, subjects in (
            ("single", ["林晓彤", "陈默"]),
            ("two_shot", ["陈默"]),
            ("group", ["林晓彤", "陈默"]),
        ):
            changed = copy.deepcopy(self.example)
            changed["shots"][0]["camera"]["framing_mode"] = mode
            changed["shots"][0]["camera"]["primary_subjects"] = subjects
            self.assertIn("FRAMING_CARDINALITY_INVALID", self._formal_codes(changed), mode)

        shoulder = copy.deepcopy(self.example)
        camera = shoulder["shots"][0]["camera"]
        camera["framing_mode"] = "over_shoulder"
        camera["primary_subjects"] = ["陈默"]
        camera["foreground_subjects"] = ["陈默"]
        self.assertIn("OVER_SHOULDER_SUBJECTS_INVALID", self._formal_codes(shoulder))

        insert = copy.deepcopy(self.example)
        shot = insert["shots"][0]
        shot["camera"]["framing_mode"] = "insert"
        shot["viewpoint"]["owner_type"] = "relationship"
        shot["viewpoint"]["reading_priority"] = "relationship"
        self.assertIn("INSERT_VIEWPOINT_INVALID", self._formal_codes(insert))

    def test_camera_reasons_and_follow_response_must_be_realized(self) -> None:
        missing = copy.deepcopy(self.example)
        missing["shots"][0]["camera"]["shot_size_reason"] = ""
        missing["shots"][0]["camera"]["angle_reason"] = ""
        self.assertIn("CAMERA_DESIGN_UNRESOLVED", self._formal_codes(missing))

        follow = copy.deepcopy(self.example)
        follow["shots"][0]["viewpoint"]["camera_response"] = "follow"
        follow["shots"][0]["camera"]["movement"] = {
            "type": "fixed",
            "trigger": "",
            "speed": "",
            "path": "",
            "end_condition": "",
            "reason": "人物没有移动，摄影机保持原位。",
        }
        self.assertIn("CAMERA_RESPONSE_NOT_REALIZED", self._formal_codes(follow))

        for response in ("reframe", "reveal"):
            unrealized = copy.deepcopy(self.example)
            shot = unrealized["shots"][0]
            shot["viewpoint"]["camera_response"] = response
            shot["camera"]["movement"] = {
                "type": "fixed",
                "trigger": "",
                "speed": "",
                "path": "",
                "end_condition": "",
                "reason": "摄影机保持原位，没有发生画面重构。",
            }
            shot["shot_flow"] = [
                item for item in shot["shot_flow"] if item["owner"] not in {"blocking", "focus"}
            ]
            self.assertIn("CAMERA_RESPONSE_NOT_REALIZED", self._formal_codes(unrealized), response)

    def test_viewing_design_and_formal_viewpoint_match_in_both_directions(self) -> None:
        formal_changed = copy.deepcopy(self.example)
        formal_changed["shots"][0]["viewpoint"]["reason"] += "正式层单独改变。"
        self.assertIn(
            "VIEWING_DESIGN_EXECUTION_DIVERGED",
            self._workspace_codes(copy.deepcopy(self.workspace), formal_changed),
        )

        workspace_changed = copy.deepcopy(self.workspace)
        workspace_changed["scene_strategies"][0]["topology"][0]["viewing_design"]["reason"] += (
            "Gate 2 单独改变。"
        )
        self.assertIn(
            "VIEWING_DESIGN_EXECUTION_DIVERGED",
            self._workspace_codes(workspace_changed, copy.deepcopy(self.example)),
        )

    def _uniform_camera_fixture(self, uniformity_intent: str) -> tuple[dict, dict]:
        data = copy.deepcopy(self.example)
        workspace = copy.deepcopy(self.workspace)
        by_id = {shot["shot_id"]: shot for shot in data["shots"]}
        for strategy in workspace["scene_strategies"]:
            strategy["camera_grammar"] = {
                "dominant_principle": "从厨房门口同一位置见证每个动作，不替人物改变距离。",
                "change_triggers": [],
                "progression": "只由人物动作和声音改变，摄影机评价保持不变。",
                "uniformity_intent": uniformity_intent,
            }
            for index, unit in enumerate(strategy["topology"]):
                viewing = {
                    "owner_type": "space",
                    "owner_refs": ["厨房关系空间"],
                    "visible_subjects": [],
                    "offscreen_subjects": [],
                    "reading_priority": "space",
                    "framing_intent": "space",
                    "camera_response": "observe",
                    "frame_axis": "centered",
                    "aspect_ratio_fit": f"单元 {index + 1} 在16:9中心轴保持空间评价一致。",
                    "reason": f"单元 {index + 1} 保持门口观察，以读取该单元独有的动作节拍。",
                }
                for shot_id in unit["shot_refs"]:
                    shot = by_id[shot_id]
                    viewing["visible_subjects"] = list(shot["staging"]["visible_subjects"])
                    viewing["offscreen_subjects"] = list(shot["staging"]["offscreen_subjects"])
                    shot["viewpoint"] = {
                        key: copy.deepcopy(viewing[key])
                        for key in ("owner_type", "owner_refs", "reading_priority", "camera_response", "reason")
                    }
                    shot["camera"].update(
                        {
                            "framing_mode": "space",
                            "primary_subjects": [],
                            "foreground_subjects": [],
                            "shot_size": "全景",
                            "angle": "平视",
                            "shot_size_reason": f"单元 {index + 1} 保留厨房全部动作路径。",
                            "angle_reason": f"单元 {index + 1} 不改变人物权力评价。",
                            "frame_axis": "centered",
                            "aspect_ratio_reason": viewing["aspect_ratio_fit"],
                            "movement": {
                                "type": "fixed",
                                "trigger": "",
                                "speed": "",
                                "path": "",
                                "end_condition": "",
                                "reason": f"单元 {index + 1} 让该动作在稳定空间中自行发生。",
                            },
                        }
                    )
                    shot["execution_text"] = canonical_execution_text(shot)
                unit["viewing_design"] = viewing
                for item in strategy["dialogue_edit_plan"]:
                    for step in item["picture_steps"]:
                        if step["topology_unit_id"] == unit["unit_id"]:
                            step["owner_type"] = viewing["owner_type"]
                            step["owner_refs"] = list(viewing["owner_refs"])
                            step["framing_intent"] = viewing["framing_intent"]
        return data, workspace

    def test_uniform_camera_requires_specific_intent_but_has_no_category_quota(self) -> None:
        data, workspace = self._uniform_camera_fixture("")
        self.assertIn("CAMERA_UNIFORMITY_INTENT_REQUIRED", self._workspace_codes(workspace, data))

        intent = "这场仪式要求每个人在同一正面尺度中被依次衡量；改变角度或距离会破坏第三次重复才显出的缺席。"
        data, workspace = self._uniform_camera_fixture(intent)
        codes = self._workspace_codes(workspace, data)
        self.assertEqual(set(), codes)
        self.assertNotIn("CAMERA_UNIFORMITY_INTENT_REQUIRED", codes)
        self.assertNotIn("CAMERA_DECISION_TEMPLATE_COLLAPSE", codes)

    def test_periodic_category_rotation_with_copied_reasons_still_collapses(self) -> None:
        data = copy.deepcopy(self.example)
        workspace = copy.deepcopy(self.workspace)
        topology = workspace["scene_strategies"][0]["topology"]
        self.assertGreater(len({shot["camera"]["framing_mode"] for shot in data["shots"]}), 1)
        self.assertGreater(len({shot["camera"]["shot_size"] for shot in data["shots"]}), 1)
        for unit in topology:
            unit["viewing_design"]["reason"] = "增强电影感与节奏。"
        for shot in data["shots"]:
            shot["camera"]["shot_size_reason"] = "增强电影感与节奏。"
            shot["camera"]["angle_reason"] = "增强电影感与节奏。"
            shot["camera"]["movement"]["reason"] = "增强电影感与节奏。"
            shot["execution_text"] = canonical_execution_text(shot)
        self.assertIn("CAMERA_DECISION_TEMPLATE_COLLAPSE", self._workspace_codes(workspace, data))

    def test_listener_hold_and_insert_detail_are_valid_without_forcing_both_in_every_scene(self) -> None:
        listener = copy.deepcopy(self.example)
        source_speakers = {
            item["dialogue_id"]: item["speaker"] for item in listener["source"]["dialogue_lines"]
        }
        shot = next(
            item
            for item in listener["shots"]
            if any(seg["delivery"] == "onscreen" for seg in item["sound"]["dialogue_segments"])
        )
        segment = next(seg for seg in shot["sound"]["dialogue_segments"] if seg["delivery"] == "onscreen")
        speaker = source_speakers[segment["dialogue_id"]]
        listener_subject = next(name for name in ("林晓彤", "陈默") if name != speaker)
        segment["delivery"] = "os"
        shot["viewpoint"] = {
            "owner_type": "subject",
            "owner_refs": [listener_subject],
            "reading_priority": "face",
            "camera_response": "isolate",
            "reason": "说话者继续发声，画面停在倾听者接收信息后的迟到反应。",
        }
        shot["camera"].update(
            {
                "framing_mode": "single",
                "primary_subjects": [listener_subject],
                "foreground_subjects": [],
                "shot_size_reason": "读取倾听者没有立即回答的面部反应。",
                "angle_reason": "与倾听者同高，不替他的反应下判断。",
            }
        )
        shot["staging"]["subjects"] = [listener_subject, speaker]
        shot["staging"]["visible_subjects"] = [listener_subject]
        shot["staging"]["offscreen_subjects"] = [speaker]
        listener_codes = self._formal_codes(listener)
        self.assertEqual(set(), listener_codes)
        self.assertNotIn("DIALOGUE_VISIBILITY_MISMATCH", listener_codes)
        self.assertNotIn("FRAMING_CARDINALITY_INVALID", listener_codes)

        insert = copy.deepcopy(self.example)
        insert_shot = next(shot for shot in insert["shots"] if shot["camera"]["framing_mode"] == "insert")
        insert_codes = self._formal_codes(insert)
        self.assertEqual("detail", insert_shot["viewpoint"]["reading_priority"])
        self.assertIn(insert_shot["viewpoint"]["owner_type"], {"object", "subject"})
        self.assertNotIn("INSERT_VIEWPOINT_INVALID", insert_codes)


@unittest.skipIf(OPENPYXL_ERROR is not None, "openpyxl unavailable")
class XlsxFrontendProjectionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.example = json.loads(
            (ROOT / "examples" / "kitchen-farewell-shot-data.json").read_text(encoding="utf-8")
        )
        cls.workspace = json.loads(
            (ROOT / "examples" / "kitchen-farewell-director-workspace.json").read_text(encoding="utf-8")
        )
        cls.concept = json.loads(
            (ROOT / "examples" / "unknown-room-awakening-shot-data.json").read_text(encoding="utf-8")
        )

    @staticmethod
    def _fifth_column(workbook) -> list[str]:
        sheet = workbook["导演分镜"]
        return [str(sheet.cell(row=row, column=5).value or "") for row in range(5, sheet.max_row + 1)]

    def test_xlsx_fifth_column_is_two_layer_frontend_not_backend_execution_text(self) -> None:
        dialogue_index = build_dialogue_index(self.example)
        workbook = build_openpyxl_workbook(self.example)
        projected = self._fifth_column(workbook)
        self.assertEqual(len(self.example["shots"]), len(projected))
        for shot, frontend in zip(self.example["shots"], projected):
            self.assertEqual(render_xlsx_execution_text(shot, dialogue_index), frontend)
            self.assertNotEqual(shot["execution_text"], frontend)
            self.assertRegex(frontend, FRONTEND_RE)
            self.assertEqual(1, frontend.count("\n"))

        joined = "\n".join(projected)
        for forbidden in FORBIDDEN_FRONTEND_TEXT:
            self.assertNotIn(forbidden, joined)
        self.assertIsNone(MACHINE_ID_RE.search(joined))

    def test_frontend_dialogue_is_verbatim_once_and_in_source_order(self) -> None:
        joined = "\n".join(self._fifth_column(build_openpyxl_workbook(self.example)))
        previous = -1
        for line in self.example["source"]["dialogue_lines"]:
            text = line["text"]
            self.assertEqual(1, joined.count(text), text)
            position = joined.index(text)
            self.assertGreater(position, previous, text)
            previous = position
            self.assertIn(line["speaker"], joined)

    def test_xlsx_summary_and_complete_two_sheet_parity_metrics(self) -> None:
        projection = project_xlsx(self.example)
        workbook = build_openpyxl_workbook(self.example)
        self.assertEqual(projection["summary"], workbook["导演分镜"]["A2"].value)
        self.assertNotIn("contract", workbook["导演分镜"]["A2"].value)
        self.assertNotIn("project", workbook["导演分镜"]["A2"].value)

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "storyboard.xlsx"
            export_xlsx(self.example, output)
            verification = verify_xlsx(self.example, output)
        self.assertEqual("PASS", verification["status"])
        self.assertEqual({"导演分镜": "PASS", "导演设计": "PASS"}, verification["sheet_parity"])
        self.assertEqual(projection_sha256(self.example), verification["projection_sha256"])
        expected_dialogue_count = sum(
            len(shot["sound"]["dialogue_segments"]) for shot in self.example["shots"]
        )
        self.assertEqual(expected_dialogue_count, verification["dialogue_segment_count"])
        self.assertEqual(10, verification["director_design_field_count"])

    def test_director_design_keeps_ten_summary_rows_and_human_assumptions(self) -> None:
        workbook = build_openpyxl_workbook(self.concept)
        sheet = workbook["导演设计"]
        labels = [sheet.cell(row=row, column=1).value for row in range(4, 14)]
        self.assertEqual(
            [
                "场景任务",
                "戏剧问题",
                "转折点",
                "观众位置",
                "视点策略",
                "情绪弧线",
                "人物调度",
                "摄影策略",
                "声音策略",
                "节奏策略",
            ],
            labels,
        )
        values = "\n".join(
            str(cell.value)
            for row in sheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        for item in self.concept["assumptions"]:
            self.assertIn(item["statement"], values)
            self.assertIn(item["impact"], values)
            self.assertNotIn(item["assumption_id"], values)
            self.assertNotIn(f"· {item['status']}", values)
        self.assertIn("待确认项 1", values)

    def test_build_keeps_agent_backend_execution_text_and_xlsx_projection_separate(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "delivery"
            report, paths = build_outputs(self.example, self.workspace, output)
            self.assertEqual("READY", report["status"])
            built = json.loads(paths["json"].read_text(encoding="utf-8"))
            markdown = paths["markdown"].read_text(encoding="utf-8")
            from openpyxl import load_workbook

            workbook = load_workbook(paths["xlsx"], read_only=False, data_only=True)
            try:
                frontend = self._fifth_column(workbook)
            finally:
                workbook.close()
            for index, shot in enumerate(self.example["shots"]):
                self.assertEqual(shot["execution_text"], built["shots"][index]["execution_text"])
                self.assertIn(markdown_cell(shot["execution_text"]), markdown)
                self.assertNotEqual(shot["execution_text"], frontend[index])
            self.assertIn("情绪弧线", markdown)
            self.assertIn(self.example["director_design"]["emotional_arc"], markdown)
            verification = report["artifacts"]["storyboard_xlsx"]["verification"]
            self.assertEqual("PASS", verification["status"])
            self.assertEqual({"导演分镜": "PASS", "导演设计": "PASS"}, verification["sheet_parity"])
            self.assertEqual(projection_sha256(built), verification["projection_sha256"])


if __name__ == "__main__":
    unittest.main()
